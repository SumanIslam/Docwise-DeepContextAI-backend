"""
Docwise — answer-quality evaluation (step 2 of 2: score the graded workbook).

Reads the grading.xlsx you filled in and reports answer accuracy and
groundedness, overall and by document type, query language and question type.

Usage (from the backend folder)
-------------------------------
    python score_answers.py --workbook evaluation/answers/<run>/grading.xlsx
"""

import argparse
import json
import os
import sys
from collections import defaultdict


def load_rows(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Grading"]
    head = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(head, r)))
    return rows


def norm_correct(v):
    if v is None or str(v).strip() == "":
        return None
    s = str(v).strip().upper()
    if s == "NA":
        return "NA"
    try:
        f = float(s)
    except ValueError:
        return None
    return f if f in (0.0, 0.5, 1.0) else None


def summarise(rows):
    graded = [r for r in rows if isinstance(r["_correct"], float)]
    if not graded:
        return None
    n = len(graded)
    return {
        "n": n,
        "accuracy": round(sum(r["_correct"] for r in graded) / n, 3),
        "fully_correct": round(sum(r["_correct"] == 1.0 for r in graded) / n, 3),
        "partly_correct": round(sum(r["_correct"] == 0.5 for r in graded) / n, 3),
        "wrong": round(sum(r["_correct"] == 0.0 for r in graded) / n, 3),
        "grounded_yes": round(sum(str(r.get("grounded", "")).strip().lower() == "yes" for r in graded) / n, 3),
        "grounded_partial": round(sum(str(r.get("grounded", "")).strip().lower() == "partial" for r in graded) / n, 3),
        "hallucinated": round(sum(str(r.get("grounded", "")).strip().lower() == "no" for r in graded) / n, 3),
        "language_ok": round(sum(str(r.get("language_ok", "")).strip().lower() == "yes" for r in graded) / n, 3),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", default=None, help="where to write scores.json (default: next to the workbook)")
    args = ap.parse_args()

    rows = load_rows(args.workbook)
    for r in rows:
        r["_correct"] = norm_correct(r.get("answer_correct"))
    ungraded = [r["sample_id"] for r in rows if r["_correct"] is None]
    excluded = [r["sample_id"] for r in rows if r["_correct"] == "NA"]
    if ungraded:
        print(f"[S] WARNING: {len(ungraded)} rows are not graded yet: {ungraded[:10]}")

    result = {"workbook": os.path.abspath(args.workbook), "rows": len(rows),
              "ungraded": ungraded, "excluded_NA": excluded,
              "overall": summarise(rows), "groups": {}}
    for field in ["doc_type", "query_language", "question_type"]:
        buckets = defaultdict(list)
        for r in rows:
            buckets[r.get(field)].append(r)
        result["groups"][field] = {str(k): summarise(v) for k, v in sorted(buckets.items(), key=lambda x: str(x[0]))
                                   if summarise(v)}

    out = args.out or os.path.join(os.path.dirname(os.path.abspath(args.workbook)), "scores.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    o = result["overall"]
    if not o:
        sys.exit("[S] nothing graded yet.")
    print(f"\n[S] OVERALL  (n={o['n']}, {len(excluded)} excluded)")
    print(f"  accuracy (1 / 0.5 / 0 averaged) : {o['accuracy']:.3f}")
    print(f"  fully correct                   : {o['fully_correct']:.3f}")
    print(f"  partly correct                  : {o['partly_correct']:.3f}")
    print(f"  wrong                           : {o['wrong']:.3f}")
    print(f"  grounded in the context         : {o['grounded_yes']:.3f}")
    print(f"  hallucinated                    : {o['hallucinated']:.3f}")
    print(f"  answered in the right language  : {o['language_ok']:.3f}")
    for field, groups in result["groups"].items():
        print(f"\n[S] by {field}")
        for k, s in groups.items():
            print(f"  {k[:28]:<30}n={s['n']:<4}accuracy {s['accuracy']:.3f}   grounded {s['grounded_yes']:.3f}")
    print(f"\n[S] written to {out}")


if __name__ == "__main__":
    main()
